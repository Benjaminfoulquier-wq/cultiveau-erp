"""La commande groupée : les besoins de plusieurs adhérents, regroupés chez un fournisseur référencé, au tarif négocié."""

import base64
import io

from odoo import api, fields, models
from odoo.exceptions import UserError


class CommandeGroupee(models.Model):
    _name = "cultiveau.commande.groupee"
    _description = "Commande groupée réseau"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "date_limite desc, id desc"

    name = fields.Char("Commande groupée", required=True)
    fournisseur_id = fields.Many2one("res.partner", "Fournisseur", required=True, domain=[("is_company", "=", True)], tracking=True)
    date_limite = fields.Date("Besoins à déposer avant le", required=True, tracking=True)
    etat = fields.Selection([("ouverte", "Ouverte aux adhérents"), ("cloturee", "Clôturée"), ("commandee", "Commandes passées")], "État", default="ouverte", tracking=True)
    ligne_ids = fields.One2many("cultiveau.commande.groupee.ligne", "groupee_id", "Besoins")
    mes_ligne_ids = fields.One2many("cultiveau.commande.groupee.ligne", "groupee_id", "Mes besoins", domain=lambda self: [("company_id", "=", self.env.company.id)])
    total_ht = fields.Float("Total HT", compute="_compute_totaux", digits="Product Price")
    nb_adherents = fields.Integer("Adhérents", compute="_compute_totaux")
    purchase_ids = fields.One2many("purchase.order", "cultiveau_groupee_id", "Commandes d'achat")
    note = fields.Text("Conditions négociées", help="Franco, remise supplémentaire, délai, livraison : ce que le fournisseur a accordé pour cette commande.")

    @api.depends("ligne_ids.sous_total", "ligne_ids.company_id")
    def _compute_totaux(self):
        for g in self:
            g.total_ht = sum(g.ligne_ids.mapped("sous_total"))
            g.nb_adherents = len(g.ligne_ids.company_id)

    def action_cloturer(self):
        self.write({"etat": "cloturee"})
        return True

    def action_rouvrir(self):
        self.filtered(lambda g: g.etat == "cloturee").write({"etat": "ouverte"})
        return True

    def action_commander(self):
        """Une commande d'achat par adhérent, chez le fournisseur, au tarif négocié ; chaque adhérent la reçoit et la paie lui-même."""
        for g in self:
            if not g.ligne_ids:
                raise UserError("Aucun besoin déposé.")
            for societe, lignes in g.ligne_ids.grouped("company_id").items():
                lignes = lignes.filtered(lambda l: not l.purchase_line_id)
                if not lignes:
                    continue
                commande = self.env["purchase.order"].with_company(societe).sudo().create({
                    "partner_id": g.fournisseur_id.id, "company_id": societe.id, "origin": g.name, "cultiveau_groupee_id": g.id,
                    "notes": g.note or False,
                    "order_line": [(0, 0, {"product_id": l.product_id.id, "name": l.product_id.display_name, "product_qty": l.qty, "price_unit": l.prix,
                                           "product_uom": l.product_id.uom_po_id.id or l.product_id.uom_id.id, "date_planned": fields.Datetime.now()}) for l in lignes]})
                for l, pl in zip(lignes, commande.order_line):
                    l.purchase_line_id = pl
            g.etat = "commandee"
            g.message_post(body=f"{len(g.purchase_ids)} commande(s) d'achat créée(s) chez {g.fournisseur_id.name}, total {g.total_ht:,.2f} € HT.".replace(",", " "))
        return True

    def action_recapitulatif(self):
        """Le tableau pour le fournisseur : par article, la quantité totale et la part de chaque adhérent."""
        self.ensure_one()
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Commande groupée"
        societes = self.ligne_ids.company_id.sorted("name")
        ws.append([f"Commande groupée Cultiveau — {self.fournisseur_id.name} — {self.name}"])
        ws.append([])
        entete = ["Réf. fournisseur", "Référence", "Article", "Unité", "Quantité totale", "Prix net HT", "Total HT"] + [s.name for s in societes]
        ws.append(entete)
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183840")
        for produit, lignes in self.ligne_ids.grouped("product_id").items():
            info = produit.seller_ids.filtered(lambda s: s.partner_id == self.fournisseur_id)[:1]
            total = sum(lignes.mapped("qty"))
            prix = lignes[0].prix
            ws.append([info.product_code or "", produit.default_code or "", produit.name, produit.uom_po_id.name or produit.uom_id.name, total, prix, round(total * prix, 2)]
                      + [sum(lignes.filtered(lambda l: l.company_id == s).mapped("qty")) or "" for s in societes])
        ws.append([])
        ws.append(["", "", "", "", "", "Total HT", round(self.total_ht, 2)])
        for i, largeur in enumerate([18, 16, 44, 8, 14, 12, 14] + [18] * len(societes), start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largeur
        sortie = io.BytesIO()
        wb.save(sortie)
        piece = self.env["ir.attachment"].create({"name": f"commande-groupee-{self.fournisseur_id.name}.xlsx".replace(" ", "-"), "datas": base64.b64encode(sortie.getvalue()),
                                                  "res_model": self._name, "res_id": self.id, "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        return {"type": "ir.actions.act_url", "url": f"/web/content/{piece.id}?download=true", "target": "self"}


class CommandeGroupeeLigne(models.Model):
    _name = "cultiveau.commande.groupee.ligne"
    _description = "Besoin d'un adhérent dans une commande groupée"

    groupee_id = fields.Many2one("cultiveau.commande.groupee", required=True, ondelete="cascade")
    company_id = fields.Many2one("res.company", "Adhérent", required=True, default=lambda self: self.env.company)
    product_id = fields.Many2one("product.product", "Article", required=True, domain=[("purchase_ok", "=", True)])
    qty = fields.Float("Quantité", digits="Product Unit of Measure", default=1.0, required=True)
    prix = fields.Float("Prix net HT", digits="Product Price", compute="_compute_prix", store=True, readonly=False)
    sous_total = fields.Float("Total HT", compute="_compute_sous_total", store=True, digits="Product Price")
    note = fields.Char("Précision")
    purchase_line_id = fields.Many2one("purchase.order.line", "Ligne de commande", readonly=True)

    @api.depends("product_id", "groupee_id.fournisseur_id")
    def _compute_prix(self):
        for l in self:
            info = l.product_id.seller_ids.filtered(lambda s: s.partner_id == l.groupee_id.fournisseur_id)[:1] if l.product_id else None
            l.prix = info.price if info else (l.product_id.with_company(l.company_id).standard_price if l.product_id else 0.0)

    @api.depends("qty", "prix")
    def _compute_sous_total(self):
        for l in self:
            l.sous_total = l.qty * l.prix


class PurchaseOrder(models.Model):
    _inherit = "purchase.order"

    cultiveau_groupee_id = fields.Many2one("cultiveau.commande.groupee", "Commande groupée", readonly=True, index=True)
