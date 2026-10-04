import base64
import io

from markupsafe import Markup, escape

from odoo import fields, models
from odoo.exceptions import UserError


class ImportArticles(models.TransientModel):
    _name = "cultiveau.import.articles"
    _description = "Importer mes articles (Excel, CSV)"

    fichier = fields.Binary("Fichier", required=True, help="Excel (.xlsx, .xls) ou CSV : l'export de votre ancien logiciel ou un tarif fournisseur. "
                                                           "Les colonnes sont reconnues par leur intitulé.")
    nom_fichier = fields.Char()
    fournisseur_id = fields.Many2one("res.partner", "Fournisseur par défaut", domain=[("is_company", "=", True)],
                                     help="Quand le fichier n'a pas de colonne « Fournisseur » ou « Marque ».")
    categorie_id = fields.Many2one("product.category", "Catégorie par défaut", help="Quand le fichier n'a pas de colonne « Famille » ou « Catégorie ».")
    mettre_a_jour = fields.Boolean("Mettre à jour les articles déjà présents", default=True,
                                   help="Même référence (ou même code-barres) : l'article est complété ; sinon il est laissé tel quel.")
    apercu = fields.Html("Aperçu", readonly=True, sanitize=False)
    resultat = fields.Html("Résultat", readonly=True, sanitize=False)

    def _rouvrir(self):
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "new"}

    def action_analyser(self):
        self.ensure_one()
        try:
            lu = self.env["product.template"].cultiveau_analyser_articles(self.nom_fichier, base64.b64decode(self.fichier))
        except ValueError as e:
            raise UserError(f"Fichier non reconnu : {e}.") from e
        corr, en_tete = lu["correspondance"], lu["en_tete"]
        libelles = {"reference": "Référence", "nom": "Désignation", "description": "Description", "fournisseur": "Fournisseur", "ref_fournisseur": "Réf. fournisseur",
                    "categorie": "Catégorie", "sous_categorie": "Sous-catégorie", "prix": "Prix de vente HT", "prix_achat": "Prix d'achat HT", "unite": "Unité",
                    "tva": "TVA", "code_barre": "Code-barres", "stock": "Stock (non importé : à saisir en inventaire)", "poids": "Poids", "conditionnement": "Conditionnement",
                    "dn": "DN", "pn": "PN", "matiere": "Matière", "raccordement": "Raccordement", "fiche": "Fiche technique"}
        reconnues = "".join(f"<li><b>{escape(str(en_tete[i]))}</b> → {escape(libelles.get(c, c))}</li>" for c, i in sorted(corr.items(), key=lambda x: x[1]))
        inconnues = ", ".join(escape(x) for x in lu["inconnus"]) or "aucune"
        apercu = "".join("<tr>" + "".join(f"<td>{escape(str(v) if v is not None else '')}</td>" for v in l[:8]) + "</tr>" for l in lu["lignes"][:3])
        self.apercu = Markup(f"<p><b>{len(lu['lignes'])} lignes</b> à importer.</p><p>Colonnes reconnues :</p><ul>{reconnues}</ul>"
                             f"<p>Colonnes laissées de côté : {inconnues}.</p><table class='table table-sm'>{apercu}</table>")
        return self._rouvrir()

    def action_importer(self):
        self.ensure_one()
        try:
            bilan = self.env["product.template"].cultiveau_importer_articles(self.nom_fichier, base64.b64decode(self.fichier), self.fournisseur_id,
                                                                            self.categorie_id, self.mettre_a_jour)
        except ValueError as e:
            raise UserError(f"Fichier non reconnu : {e}.") from e
        texte = f"<p><b>{bilan['crees']} articles créés</b>, {bilan['maj']} mis à jour, {bilan['ignores']} ignorés.</p>"
        if bilan["stock_ignore"]:
            texte += "<p>La colonne de stock n'est pas importée ici : les quantités se saisissent dans Inventaire → Ajustements, une fois les articles créés.</p>"
        if bilan["erreurs"]:
            texte += "<p>À regarder :</p><ul>" + "".join(f"<li>{escape(e)}</li>" for e in bilan["erreurs"][:50]) + "</ul>"
        self.resultat = Markup(texte)
        self.apercu = False
        return self._rouvrir()

    def action_modele(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Articles"
        colonnes = ["Référence", "Désignation", "Fournisseur / marque", "Réf. fournisseur", "Famille", "Sous-famille", "Prix de vente HT", "Prix d'achat HT",
                    "TVA %", "Unité", "Code-barres", "Conditionnement", "Poids (kg)", "DN", "PN", "Matière", "Raccordement", "Description"]
        ws.append(colonnes)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183840")
        ws.append(["TUB-PE-32 (exemple)", "Tube PE 32 PN10", "ATUSA", "PE32-10", "Tuyaux", "PE", "1,45", "0,98", "20", "mètre", "", "Rouleau 100 m", "0,21", "32", "10", "PE", "compression", ""])
        for i, largeur in enumerate([20, 36, 22, 16, 18, 16, 16, 16, 8, 10, 16, 18, 10, 8, 8, 14, 16, 40], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largeur
        sortie = io.BytesIO()
        wb.save(sortie)
        piece = self.env["ir.attachment"].create({"name": "articles-cultiveau.xlsx", "datas": base64.b64encode(sortie.getvalue()),
                                                  "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        return {"type": "ir.actions.act_url", "url": f"/web/content/{piece.id}?download=true", "target": "self"}
