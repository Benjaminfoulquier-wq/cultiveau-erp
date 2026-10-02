import base64
import io
import json

from odoo import fields, models
from odoo.exceptions import UserError


class ImportMatrice(models.TransientModel):
    _name = "cultiveau.catalogue.import.matrice"
    _description = "Importer la matrice de produits Cultiveau (Excel)"

    fichier = fields.Binary("Classeur .xlsx", required=True)
    nom_fichier = fields.Char()
    fournisseur_id = fields.Many2one("res.partner", "Fournisseur par défaut", domain=[("is_company", "=", True)],
                                     help="Utilisé quand la colonne Fournisseur est vide.")
    resultat = fields.Text("Résultat", readonly=True)

    def action_importer(self):
        self.ensure_one()
        contenu = base64.b64decode(self.fichier)
        try:
            bilan = self.env["product.template"].cultiveau_importer_matrice(contenu, self.fournisseur_id)
        except Exception as e:  # classeur illisible
            raise UserError(f"Impossible de lire le classeur : {e}")
        texte = f"Articles créés : {bilan['crees']} ; mis à jour : {bilan['maj']} ; ignorés : {bilan['ignores']}."
        if bilan["erreurs"]:
            texte += "\n\nÀ corriger :\n- " + "\n- ".join(bilan["erreurs"][:50])
        self.resultat = texte
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "new"}

    def action_modele(self):
        """Le classeur modèle vide, à remplir."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Produits"
        ws.append(["MATRICE D'IMPORTATION DES PRODUITS — CULTIVEAU"])
        ws.append(["Remplissez ce fichier puis importez-le : Cultiveau → Catalogue → Importer la matrice. Les colonnes * sont obligatoires."])
        ws.append([])
        ws.append(["Référence *", "Nom du produit *", "Fournisseur *", "Catégorie", "Description", "Prix HT (€)", "Unité", "Conditionnement",
                   "Poids (kg)", "URL Image", "Fiche technique", "Délai (jours)", "DN", "PN", "Matière", "Raccordement"])
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D3A52")
        ws.append(["Code unique du produit (ex: REF-001)", "Nom complet", "Nom exact du fournisseur", "Famille / sous-famille", "",
                   "", "unité, mètre, kg, litre…", "Ex. Rouleau 500 m", "", "", "Lien PDF", "", "", "", "", ""])
        sortie = io.BytesIO()
        wb.save(sortie)
        piece = self.env["ir.attachment"].create({"name": "matrice-import-produits-cultiveau.xlsx", "datas": base64.b64encode(sortie.getvalue()),
                                                  "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        return {"type": "ir.actions.act_url", "url": f"/web/content/{piece.id}?download=true", "target": "self"}


class ImportCatalogue3D(models.TransientModel):
    _name = "cultiveau.catalogue.import.3d"
    _description = "Importer le catalogue 3D (catalogue.json)"

    fichier = fields.Binary("catalogue.json", required=True)
    nom_fichier = fields.Char()
    resultat = fields.Text("Résultat", readonly=True)

    def action_importer(self):
        self.ensure_one()
        try:
            donnees = json.loads(base64.b64decode(self.fichier).decode("utf-8"))
            dicts, fiches, produits = donnees["dicts"], donnees.get("fiches", {}), donnees["produits"]
        except (ValueError, KeyError) as e:
            raise UserError(f"Ce fichier n'est pas un catalogue 3D (clés dicts, fiches, produits) : {e}")
        Produit = self.env["product.template"]
        total = {"crees": 0, "maj": 0}
        for i in range(0, len(produits), 2000):
            bilan = Produit.cultiveau_importer_3d(dicts, fiches, produits[i:i + 2000])
            total["crees"] += bilan["crees"]
            total["maj"] += bilan["maj"]
        self.resultat = f"{len(produits)} produits lus : {total['crees']} créés, {total['maj']} mis à jour."
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "new"}
