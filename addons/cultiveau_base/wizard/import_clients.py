import base64
import io

from markupsafe import Markup, escape

from odoo import fields, models
from odoo.exceptions import UserError

from ..models.res_partner import TYPES


class ImportClients(models.TransientModel):
    _name = "cultiveau.import.clients"
    _description = "Importer mes clients (Excel, CSV)"

    fichier = fields.Binary("Fichier", required=True, help="Excel (.xlsx, .xls) ou CSV : l'export de votre ancien logiciel, de votre comptable, "
                                                           "de l'assistant téléphonique ou de votre téléphone. Les colonnes sont reconnues par leur intitulé.")
    nom_fichier = fields.Char()
    type_defaut = fields.Selection(TYPES, "Ce sont des", default="agriculteur", required=True,
                                   help="Le type donné aux contacts quand le fichier n'a pas de colonne « Type ».")
    mettre_a_jour = fields.Boolean("Mettre à jour les contacts déjà connus", default=True,
                                   help="Un contact au même numéro, e-mail, ou nom et code postal est complété ; sinon il est laissé tel quel.")
    apercu = fields.Html("Aperçu", readonly=True, sanitize=False)
    resultat = fields.Html("Résultat", readonly=True, sanitize=False)

    def _lire(self):
        try:
            return self.env["cultiveau.import.clients.moteur"].analyser(self.nom_fichier, base64.b64decode(self.fichier))
        except ValueError as e:
            raise UserError(f"Fichier non reconnu : {e}.") from e

    def _rouvrir(self):
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "new"}

    def action_analyser(self):
        """Montre les colonnes reconnues et les premières lignes avant d'importer."""
        self.ensure_one()
        lu = self._lire()
        corr, en_tete = lu["correspondance"], lu["en_tete"]
        libelles = {"nom": "Nom", "prenom": "Prénom", "civilite": "Civilité", "exploitation": "Exploitation / société", "telephone": "Téléphone",
                    "telephone2": "Autre téléphone", "email": "E-mail", "adresse": "Adresse", "adresse2": "Complément", "code_postal": "Code postal",
                    "commune": "Commune", "pays": "Pays", "siret": "SIRET", "tva": "TVA", "type": "Type", "surface": "Surface", "cultures": "Cultures",
                    "materiel": "Matériel", "produits": "Produits", "commercial": "Commercial", "notes": "Notes"}
        reconnues = "".join(f"<li><b>{escape(str(en_tete[i]))}</b> → {escape(libelles.get(c, c))}</li>" for c, i in sorted(corr.items(), key=lambda x: x[1]))
        inconnues = ", ".join(escape(x) for x in lu["inconnus"]) or "aucune"
        apercu = "".join("<tr>" + "".join(f"<td>{escape(str(v) if v is not None else '')}</td>" for v in l[:8]) + "</tr>" for l in lu["lignes"][:3])
        self.apercu = Markup(f"<p><b>{len(lu['lignes'])} lignes</b> à importer.</p><p>Colonnes reconnues :</p><ul>{reconnues}</ul>"
                             f"<p>Colonnes laissées de côté : {inconnues}.</p><table class='table table-sm'>{apercu}</table>")
        return self._rouvrir()

    def action_importer(self):
        self.ensure_one()
        try:
            bilan = self.env["cultiveau.import.clients.moteur"].importer(self.nom_fichier, base64.b64decode(self.fichier), self.type_defaut,
                                                                        mettre_a_jour=self.mettre_a_jour)
        except ValueError as e:
            raise UserError(f"Fichier non reconnu : {e}.") from e
        texte = f"<p><b>{bilan['crees']} contacts créés</b>, {bilan['maj']} mis à jour, {bilan['ignores']} ignorés (sans nom, ou déjà connus).</p>"
        if bilan["erreurs"]:
            texte += "<p>À regarder :</p><ul>" + "".join(f"<li>{escape(e)}</li>" for e in bilan["erreurs"][:50]) + "</ul>"
        self.resultat = Markup(texte)
        self.apercu = False
        return self._rouvrir()

    def action_modele(self):
        """Le classeur modèle, à remplir."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clients"
        colonnes = ["Nom", "Prénom", "Exploitation", "Téléphone portable", "Téléphone fixe", "E-mail", "Adresse", "Code postal", "Commune",
                    "Type (client, prospect, fournisseur, partenaire)", "Surface (ha)", "Cultures", "Matériel installé", "Commercial", "Notes"]
        ws.append(colonnes)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183840")
        ws.append(["Durand (exemple)", "Paul", "EARL des Oliviers", "06 12 34 56 78", "", "paul.durand@exemple.fr", "12 chemin des Prés", "30100", "Alès",
                   "client", "35", "vigne, oliviers", "Goutte-à-goutte 12 ha, pompe 15 m³/h", "", "Rappeler avant la taille"])
        for i, largeur in enumerate([22, 14, 24, 18, 16, 28, 28, 12, 18, 30, 12, 24, 36, 16, 30], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largeur
        sortie = io.BytesIO()
        wb.save(sortie)
        piece = self.env["ir.attachment"].create({"name": "clients-cultiveau.xlsx", "datas": base64.b64encode(sortie.getvalue()),
                                                  "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        return {"type": "ir.actions.act_url", "url": f"/web/content/{piece.id}?download=true", "target": "self"}
